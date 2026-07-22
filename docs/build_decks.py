#!/usr/bin/env python3
"""
Build two Project Pulse (GenMedia SxS) decks in Google brand style:
  1. User Guide  — how to rate on the platform (for all evaluators)
  2. Product Overview — why / what it solves (for leadership)

Google fonts (Google Sans / Roboto / Roboto Mono) + Google brand colors.
Uses live-captured screenshots in docs/deck_assets/.
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "deck_assets")

# ---- Google brand palette ---------------------------------------------------
BLUE   = RGBColor(0x42, 0x85, 0xF4)
RED    = RGBColor(0xEA, 0x43, 0x35)
YELLOW = RGBColor(0xFB, 0xBC, 0x04)
GREEN  = RGBColor(0x34, 0xA8, 0x53)
G900   = RGBColor(0x20, 0x21, 0x24)   # near-black text
G700   = RGBColor(0x5F, 0x63, 0x68)   # secondary text
G500   = RGBColor(0x80, 0x86, 0x8B)
G200   = RGBColor(0xE8, 0xEA, 0xED)   # hairlines
G100   = RGBColor(0xF1, 0xF3, 0xF4)   # code / panel bg
G050   = RGBColor(0xF8, 0xF9, 0xFA)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
INK    = RGBColor(0x0B, 0x0E, 0x14)   # deep ink for dividers

HEAD = "Google Sans"
BODY = "Roboto"
MONO = "Roboto Mono"

SW, SH = Inches(13.333), Inches(7.5)
FOUR = [BLUE, RED, YELLOW, GREEN]


def new_deck():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    return prs


def blank(prs, bg=WHITE):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    f = s.background.fill
    f.solid(); f.fore_color.rgb = bg
    return s


def rect(s, l, t, w, h, color, line=None, line_w=None, shape=MSO_SHAPE.RECTANGLE):
    sp = s.shapes.add_shape(shape, l, t, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = line_w or Pt(1)
    sp.shadow.inherit = False
    return sp


def four_bar(s, t=Inches(0), h=Inches(0.09), l=Inches(0), total=SW):
    seg = Emu(int(total / 4))
    for i, c in enumerate(FOUR):
        rect(s, Emu(int(l) + i * int(seg)), t, seg, h, c)


def text(s, l, t, w, h, runs, size=18, bold=False, color=G900, font=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spc=1.0, wrap=True):
    tb = s.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    if isinstance(runs, str):
        runs = [(runs, {})]
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spc
    for txt, ov in runs:
        r = p.add_run(); r.text = txt
        r.font.size = Pt(ov.get("size", size))
        r.font.bold = ov.get("bold", bold)
        r.font.name = ov.get("font", font)
        r.font.color.rgb = ov.get("color", color)
    return tb, tf


def para(tf, runs, size=16, color=G700, font=BODY, bold=False, space_before=8,
         align=PP_ALIGN.LEFT, line_spc=1.05, bullet=None, bullet_color=None):
    p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.line_spacing = line_spc
    if bullet is not None:
        rb = p.add_run(); rb.text = bullet + "  "
        rb.font.size = Pt(size); rb.font.bold = True; rb.font.name = font
        rb.font.color.rgb = bullet_color or BLUE
    if isinstance(runs, str):
        runs = [(runs, {})]
    for txt, ov in runs:
        r = p.add_run(); r.text = txt
        r.font.size = Pt(ov.get("size", size))
        r.font.bold = ov.get("bold", bold)
        r.font.name = ov.get("font", font)
        r.font.color.rgb = ov.get("color", color)
    return p


def imsize(path):
    with Image.open(path) as im:
        return im.size


def framed_image(s, path, box, pad=Inches(0.06)):
    """Contain image inside box (l,t,w,h), centered, with a hairline frame + soft shadow."""
    L, T, W, H = box
    w_px, h_px = imsize(path)
    a = w_px / h_px
    boxA = W / H
    if a > boxA:
        dw = W; dh = Emu(int(W / a))
    else:
        dh = H; dw = Emu(int(H * a))
    dl = Emu(int(L) + (int(W) - int(dw)) // 2)
    dt = Emu(int(T) + (int(H) - int(dh)) // 2)
    # soft shadow
    sh = rect(s, Emu(int(dl) + Emu(0.05*914400)), Emu(int(dt) + Emu(0.06*914400)),
              dw, dh, G200)
    # frame
    fr = rect(s, Emu(int(dl) - int(pad)), Emu(int(dt) - int(pad)),
              Emu(int(dw) + 2*int(pad)), Emu(int(dh) + 2*int(pad)), WHITE,
              line=G200, line_w=Pt(1))
    pic = s.shapes.add_picture(path, dl, dt, dw, dh)
    pic.line.color.rgb = G200; pic.line.width = Pt(0.75)
    return dl, dt, dw, dh


def kicker(s, l, t, txt, color=BLUE):
    rect(s, l, Emu(int(t) + Emu(0.02*914400)), Inches(0.16), Inches(0.30), color)
    text(s, Emu(int(l) + Emu(0.30*914400)), t, Inches(8), Inches(0.34),
         txt.upper(), size=12.5, bold=True, color=color, font=HEAD)


def footer(s, name, idx):
    text(s, Inches(0.6), Inches(7.06), Inches(8), Inches(0.3),
         name, size=9, color=G500, font=BODY)
    text(s, Inches(11.4), Inches(7.06), Inches(1.33), Inches(0.3),
         str(idx), size=9, color=G500, font=BODY, align=PP_ALIGN.RIGHT)
    rect(s, Inches(0.6), Inches(7.02), Inches(12.13), Pt(0.75), G200)


# ---------------------------------------------------------------------------
# Reusable slide templates
# ---------------------------------------------------------------------------
def title_slide(prs, eyebrow, title_runs, subtitle, meta):
    s = blank(prs)
    four_bar(s)
    rect(s, Inches(0.9), Inches(1.35), Inches(0.9), Inches(0.16), BLUE)
    text(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(0.5),
         eyebrow.upper(), size=15, bold=True, color=G700, font=HEAD)
    text(s, Inches(0.88), Inches(2.25), Inches(11.6), Inches(2.2),
         title_runs, size=54, bold=True, color=G900, font=HEAD, line_spc=1.0)
    text(s, Inches(0.9), Inches(4.5), Inches(10.8), Inches(0.9),
         subtitle, size=20, color=G700, font=BODY, line_spc=1.15)
    # colored dots
    for i, c in enumerate(FOUR):
        rect(s, Emu(int(Inches(0.92)) + i*int(Inches(0.34))), Inches(5.7),
             Inches(0.2), Inches(0.2), c, shape=MSO_SHAPE.OVAL)
    text(s, Inches(0.9), Inches(6.7), Inches(11.5), Inches(0.4),
         meta, size=12.5, color=G500, font=BODY)
    return s


def section_slide(prs, num, title, sub):
    s = blank(prs, INK)
    four_bar(s)
    text(s, Inches(0.9), Inches(2.5), Inches(3), Inches(1.2),
         num, size=64, bold=True, color=RGBColor(0x3C,0x40,0x49), font=HEAD)
    text(s, Inches(0.9), Inches(3.55), Inches(11), Inches(1.0),
         title, size=40, bold=True, color=WHITE, font=HEAD)
    text(s, Inches(0.94), Inches(4.7), Inches(10.5), Inches(0.8),
         sub, size=18, color=RGBColor(0xBD,0xC1,0xC6), font=BODY, line_spc=1.15)
    for i, c in enumerate(FOUR):
        rect(s, Emu(int(Inches(0.94)) + i*int(Inches(0.34))), Inches(5.7),
             Inches(0.2), Inches(0.2), c, shape=MSO_SHAPE.OVAL)
    return s


def shot_left_slide(prs, deckname, idx, kick, heading, img, intro, points,
                    kick_color=BLUE):
    """Tall screenshot on the left, narrative on the right."""
    s = blank(prs)
    four_bar(s)
    kicker(s, Inches(0.6), Inches(0.42), kick, kick_color)
    text(s, Inches(0.6), Inches(0.78), Inches(6.2), Inches(0.9),
         heading, size=27, bold=True, color=G900, font=HEAD, line_spc=1.0)
    framed_image(s, os.path.join(ASSETS, img),
                 (Inches(0.6), Inches(1.75), Inches(6.35), Inches(5.05)))
    # narrative panel
    nx = Inches(7.35)
    _, tf = text(s, nx, Inches(1.85), Inches(5.4), Inches(0.9),
                 intro, size=15, color=G700, font=BODY, line_spc=1.15)
    for i, (h, d) in enumerate(points):
        c = FOUR[i % 4]
        para(tf, [(h + "  ", {"bold": True, "color": G900, "size": 15, "font": HEAD}),
                  (d, {"color": G700, "size": 14})],
             space_before=13, bullet="●", bullet_color=c, line_spc=1.1)
    footer(s, deckname, idx)
    return s


def shot_full_slide(prs, deckname, idx, kick, heading, img, points,
                    kick_color=BLUE, img_box=None):
    """Wide screenshot: heading on top, image large, callouts on the right."""
    s = blank(prs)
    four_bar(s)
    kicker(s, Inches(0.6), Inches(0.42), kick, kick_color)
    text(s, Inches(0.6), Inches(0.78), Inches(12), Inches(0.7),
         heading, size=27, bold=True, color=G900, font=HEAD)
    box = img_box or (Inches(0.6), Inches(1.75), Inches(8.4), Inches(5.0))
    framed_image(s, os.path.join(ASSETS, img), box)
    nx = Inches(9.35)
    _, tf = text(s, nx, Inches(1.85), Inches(3.4), Inches(0.4),
                 [("What you're seeing", {"bold": True, "color": G900,
                                          "size": 15, "font": HEAD})])
    for i, (h, d) in enumerate(points):
        c = FOUR[i % 4]
        para(tf, [(h + "  ", {"bold": True, "color": G900, "size": 13.5, "font": HEAD}),
                  (d, {"color": G700, "size": 12.5})],
             space_before=12, bullet="●", bullet_color=c, line_spc=1.08)
    footer(s, deckname, idx)
    return s


def leaderboard_slide(prs, deckname, idx, kick, heading, intro, takeaways,
                      kick_color=YELLOW):
    """Wide leaderboard screenshot band + a row of takeaway cards."""
    s = blank(prs)
    four_bar(s)
    kicker(s, Inches(0.6), Inches(0.42), kick, kick_color)
    text(s, Inches(0.6), Inches(0.78), Inches(12.1), Inches(0.6),
         heading, size=27, bold=True, color=G900, font=HEAD)
    text(s, Inches(0.6), Inches(1.45), Inches(12.1), Inches(0.55),
         intro, size=14.5, color=G700, font=BODY, line_spc=1.15)
    framed_image(s, os.path.join(ASSETS, "15_top_evaluators.png"),
                 (Inches(0.6), Inches(2.15), Inches(12.13), Inches(3.0)))
    n = len(takeaways); gap = Inches(0.22); total = Inches(12.13)
    cw = Emu((int(total) - (n-1)*int(gap)) // n); x = Inches(0.6); y = Inches(5.55)
    for (h, d, c) in takeaways:
        card = rect(s, x, y, cw, Inches(1.15), G050, line=G200, line_w=Pt(0.75),
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        try: card.adjustments[0] = 0.06
        except Exception: pass
        rect(s, x, y, Inches(0.09), Inches(1.15), c)
        _, tf = text(s, Emu(int(x)+Emu(0.28*914400)), Emu(int(y)+Emu(0.14*914400)),
                     Emu(int(cw)-Emu(0.5*914400)), Inches(0.9),
                     [(h, {"bold": True, "color": G900, "size": 14.5, "font": HEAD})])
        para(tf, d, size=12.5, color=G700, space_before=4, line_spc=1.05)
        x = Emu(int(x)+int(cw)+int(gap))
    footer(s, deckname, idx)
    return s


def code_slide(prs, deckname, idx, kick, heading, subtitle, code, notes,
               kick_color=BLUE):
    s = blank(prs)
    four_bar(s)
    kicker(s, Inches(0.6), Inches(0.42), kick, kick_color)
    text(s, Inches(0.6), Inches(0.78), Inches(12), Inches(0.6),
         heading, size=27, bold=True, color=G900, font=HEAD)
    text(s, Inches(0.6), Inches(1.4), Inches(12), Inches(0.5),
         subtitle, size=14.5, color=G700, font=BODY)
    # code panel
    panel = rect(s, Inches(0.6), Inches(2.0), Inches(7.7), Inches(4.7), G100,
                 line=G200, line_w=Pt(1), shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    try:
        panel.adjustments[0] = 0.03
    except Exception:
        pass
    tb = s.shapes.add_textbox(Inches(0.85), Inches(2.2), Inches(7.25), Inches(4.35))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_top = 0
    lines = code.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.04
        r = p.add_run(); r.text = ln if ln else " "
        r.font.name = MONO; r.font.size = Pt(11.5); r.font.color.rgb = G900
    # notes column
    nx = Inches(8.65)
    text(s, nx, Inches(2.0), Inches(4.1), Inches(0.4),
         [("Fields", {"bold": True, "color": G900, "size": 15, "font": HEAD})])
    _, ntf = text(s, nx, Inches(2.45), Inches(4.1), Inches(0.3),
                  [("", {})])
    for i, (h, d) in enumerate(notes):
        c = FOUR[i % 4]
        para(ntf, [(h + "  ", {"bold": True, "color": G900, "size": 13, "font": MONO}),
                   (d, {"color": G700, "size": 12.5})],
             space_before=10, bullet="●", bullet_color=c, line_spc=1.06)
    footer(s, deckname, idx)
    return s


def bullets_slide(prs, deckname, idx, kick, heading, intro, blocks,
                  kick_color=BLUE, two_col=True):
    s = blank(prs)
    four_bar(s)
    kicker(s, Inches(0.6), Inches(0.42), kick, kick_color)
    text(s, Inches(0.6), Inches(0.78), Inches(12), Inches(0.7),
         heading, size=27, bold=True, color=G900, font=HEAD)
    if intro:
        text(s, Inches(0.6), Inches(1.5), Inches(11.8), Inches(0.6),
             intro, size=15, color=G700, font=BODY, line_spc=1.15)
    top = Inches(2.25)
    if two_col:
        colw = Inches(5.85)
        xs = [Inches(0.6), Inches(6.9)]
        half = (len(blocks) + 1) // 2
        cols = [blocks[:half], blocks[half:]]
    else:
        colw = Inches(12.1); xs = [Inches(0.6)]; cols = [blocks]
    for ci, col in enumerate(cols):
        y = top
        for (h, d, c) in col:
            card = rect(s, xs[ci], y, colw, Inches(1.15), G050,
                        line=G200, line_w=Pt(0.75),
                        shape=MSO_SHAPE.ROUNDED_RECTANGLE)
            try: card.adjustments[0] = 0.06
            except Exception: pass
            rect(s, xs[ci], y, Inches(0.09), Inches(1.15), c)
            _, tf = text(s, Emu(int(xs[ci]) + Emu(0.28*914400)),
                         Emu(int(y) + Emu(0.14*914400)), Emu(int(colw)-Emu(0.5*914400)),
                         Inches(0.9),
                         [(h, {"bold": True, "color": G900, "size": 15.5, "font": HEAD})])
            para(tf, d, size=12.8, color=G700, space_before=4, line_spc=1.06)
            y = Emu(int(y) + int(Inches(1.32)))
    footer(s, deckname, idx)
    return s


def stat_row(s, y, stats):
    n = len(stats)
    gap = Inches(0.25)
    total = Inches(12.13)
    cw = Emu((int(total) - (n-1)*int(gap)) // n)
    x = Inches(0.6)
    for i, (big, small) in enumerate(stats):
        card = rect(s, x, y, cw, Inches(1.5), G050, line=G200, line_w=Pt(0.75),
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        try: card.adjustments[0] = 0.08
        except Exception: pass
        rect(s, x, y, cw, Inches(0.09), FOUR[i % 4])
        text(s, x, Emu(int(y)+Emu(0.28*914400)), cw, Inches(0.8), big,
             size=34, bold=True, color=G900, font=HEAD, align=PP_ALIGN.CENTER)
        text(s, x, Emu(int(y)+Emu(1.02*914400)), cw, Inches(0.4), small.upper(),
             size=11, bold=True, color=G700, font=HEAD, align=PP_ALIGN.CENTER)
        x = Emu(int(x) + int(cw) + int(gap))


# ===========================================================================
# DECK 1 — USER GUIDE
# ===========================================================================
def build_user_guide():
    prs = new_deck()
    D = "GenMedia SxS · Project Pulse — User Guide"
    n = [0]
    def nx(): n[0]+=1; return n[0]

    # 1 Title
    title_slide(
        prs, "Project Pulse · GenMedia SxS",
        [("How to Rate\n", {"color": G900}),
         ("Video · Image · Speech", {"color": BLUE})],
        "A 10-minute guide to blind side-by-side rating — how to vote fairly, "
        "read the results, and create your own evaluations.",
        "Internal evaluator guide  ·  genmedia-sxs-v2  ·  Google-IAM access")

    # 2 What & how it works
    s = blank(prs); four_bar(s)
    kicker(s, Inches(0.6), Inches(0.42), "Start here")
    text(s, Inches(0.6), Inches(0.78), Inches(12), Inches(0.7),
         "What is Project Pulse — and why your vote matters", size=27, bold=True,
         color=G900, font=HEAD)
    text(s, Inches(0.6), Inches(1.55), Inches(12.1), Inches(0.9),
         "Project Pulse compares two AI models on the same prompt — side by side, with names hidden. "
         "You watch/look/listen, score each side, and pick a winner. Your blind votes become the "
         "leaderboard we use to choose the best model for customers.",
         size=15.5, color=G700, font=BODY, line_spc=1.2)
    steps = [("1 · Pick", "Choose a modality: Video, Image or TTS.", BLUE),
             ("2 · Compare", "Two outputs, A vs B, models hidden.", RED),
             ("3 · Score", "Rate each side 1–5 on a few metrics.", YELLOW),
             ("4 · Decide", "Pick the overall winner and (optionally) say why.", GREEN)]
    x = Inches(0.6); cw = Inches(2.95); gap = Inches(0.11)
    for i,(h,d,c) in enumerate(steps):
        card = rect(s, x, Inches(2.75), cw, Inches(1.9), G050, line=G200, line_w=Pt(0.75),
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        try: card.adjustments[0]=0.06
        except Exception: pass
        rect(s, x, Inches(2.75), cw, Inches(0.1), c)
        text(s, x, Emu(int(Inches(2.98))), cw, Inches(0.5), h, size=18, bold=True,
             color=G900, font=HEAD, align=PP_ALIGN.CENTER)
        text(s, Emu(int(x)+Emu(0.2*914400)), Inches(3.6),
             Emu(int(cw)-Emu(0.4*914400)), Inches(1.0), d, size=13, color=G700,
             font=BODY, align=PP_ALIGN.CENTER, line_spc=1.1)
        x = Emu(int(x)+int(cw)+int(gap))
    band = rect(s, Inches(0.6), Inches(5.05), Inches(12.13), Inches(1.15), RGBColor(0xE8,0xF0,0xFE),
                shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    try: band.adjustments[0]=0.10
    except Exception: pass
    text(s, Inches(0.95), Inches(5.22), Inches(11.4), Inches(0.9),
         [("Two rules to remember:  ", {"bold": True, "color": BLUE, "size": 15, "font": HEAD}),
          ("it's blind (you never see model names while voting), and ", {"color": G900, "size": 14.5}),
          ("Analytics unlocks after your first 10 votes.", {"bold": True, "color": G900, "size": 14.5})],
         line_spc=1.2)
    footer(s, D, nx())

    # 3 Home
    shot_full_slide(prs, D, nx(), "Step 1 · Choose what to rate",
        "The home screen — six places to work",
        "01_home.png",
        [("Video / Image / TTS Eval", "the three blind rating arenas."),
         ("AI Evals", "auto-scores from an AI judge."),
         ("Analytics", "leaderboards — unlocks at 10 votes."),
         ("Admin", "create & run new evaluations.")])

    # 4 LDAP
    shot_full_slide(prs, D, nx(), "Step 2 · Sign in",
        "Enter your LDAP to start",
        "02_image_ldap_gate.png",
        [("Your LDAP", "identifies your votes on the leaderboard."),
         ("No password", "just your Google username."),
         ("Vote counter", "tracks progress to the 10-vote unlock.")],
        img_box=(Inches(0.6), Inches(1.75), Inches(8.4), Inches(4.7)))

    # 5 Arena at a glance (image, empty)
    shot_left_slide(prs, D, nx(), "Step 3 · The arena",
        "The rating arena at a glance",
        "03_image_blind_vote.png",
        "Every arena has the same shape, so once you learn one you know them all.",
        [("Prompt", "the instruction both models were given (top)."),
         ("A vs B", "the two outputs — model names are hidden."),
         ("Metric rows", "score each side 1–5 on each metric."),
         ("Next / Skip", "get a new pair, or skip if unsure.")])

    # 6 Worked scored example
    shot_left_slide(prs, D, nx(), "Step 3 · Score each side",
        "Give each side a 1–5 on every metric",
        "04_image_scored.png",
        "Click a number for A and for B on each row. 1 = poor, 5 = excellent. "
        "Judge each side on its own — you're not forced to split them.",
        [("Score A and B", "independently, row by row."),
         ("Be consistent", "use the same bar for both sides."),
         ("Say why (optional)", "a short note helps us learn."),
         ("Then pick a winner", "A Wins · Tie · B Wins.")],
        kick_color=GREEN)

    # 7 What metrics mean
    s = blank(prs); four_bar(s)
    kicker(s, Inches(0.6), Inches(0.42), "Reference")
    text(s, Inches(0.6), Inches(0.78), Inches(12), Inches(0.6),
         "What the metrics mean", size=27, bold=True, color=G900, font=HEAD)
    text(s, Inches(0.6), Inches(1.45), Inches(12), Inches(0.5),
         "Each modality scores 5 quick metrics, 1 (poor) to 5 (excellent).",
         size=14.5, color=G700, font=BODY)
    cols = [
        ("Video", BLUE, ["Prompt Adherence", "Visual Quality", "Motion & Physics",
                          "Temporal Consistency", "Audio-Visual Sync"]),
        ("Image", RED, ["Prompt Following", "Aesthetic", "Detail / Sharpness",
                        "Artifact-Free", "Edit Fidelity (i2i)"]),
        ("TTS", GREEN, ["Naturalness", "Style Adherence", "Expressiveness",
                        "Pacing", "Pronunciation"]),
    ]
    x = Inches(0.6); cw = Inches(3.95); gap = Inches(0.14)
    for (name, c, items) in cols:
        card = rect(s, x, Inches(2.15), cw, Inches(4.4), G050, line=G200, line_w=Pt(0.75),
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        try: card.adjustments[0]=0.04
        except Exception: pass
        rect(s, x, Inches(2.15), cw, Inches(0.55), c, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        text(s, x, Inches(2.2), cw, Inches(0.45), name.upper(), size=15, bold=True,
             color=WHITE, font=HEAD, align=PP_ALIGN.CENTER)
        _, tf = text(s, Emu(int(x)+Emu(0.32*914400)), Inches(2.95),
                     Emu(int(cw)-Emu(0.6*914400)), Inches(3.4), [("", {})])
        for it in items:
            para(tf, it, size=14, color=G900, font=BODY, space_before=13,
                 bullet="—", bullet_color=c, line_spc=1.05)
        x = Emu(int(x)+int(cw)+int(gap))
    footer(s, D, nx())

    # 8 TTS
    shot_left_slide(prs, D, nx(), "Rate audio",
        "TTS — listen to A and B, then score",
        "05_tts_blind_vote.png",
        "Play both clips (headphones help). Read the script and the style note so you "
        "know what 'good' should sound like.",
        [("Play both", "use the A and B audio players."),
         ("Score 5 metrics", "naturalness, style, expressiveness, pacing, pronunciation."),
         ("Pick winner", "Clip A · Tie · Clip B, then Submit Vote."),
         ("Skip / Next Pair", "move on when you're done.")],
        kick_color=GREEN)

    # 9 Video
    shot_left_slide(prs, D, nx(), "Rate video",
        "Video — watch both variants, then score",
        "06_video_blind_vote.png",
        "Watch each clip fully at least once. A small reference image (top-left) shows what "
        "the prompt was based on.",
        [("Play both variants", "A and B, full length."),
         ("Score 5 metrics", "adherence, quality, motion, consistency, A/V sync."),
         ("Winner + why", "Variant A · Tie · Variant B."),
         ("Submit", "or Skip if you can't decide.")],
        kick_color=RED)

    # 10 Golden rules
    bullets_slide(prs, D, nx(), "Do it well",
        "Golden rules for fair, useful ratings",
        "A few habits keep the leaderboard trustworthy.",
        [("Stay blind", "Never try to guess which model is which — score what you see.", BLUE),
         ("Judge the prompt", "Reward the output that best does what was asked.", RED),
         ("Use the full scale", "1–5 means 1–5. Don't give everything a 4.", YELLOW),
         ("Skip when unsure", "A confident skip beats a random vote.", GREEN),
         ("Add a short 'why'", "One line of reasoning is gold for the team.", BLUE),
         ("Take your time", "Watch/listen fully before scoring.", RED)])

    # 11 Unlock analytics
    shot_full_slide(prs, D, nx(), "Unlock the results",
        "Cast 10 votes to unlock Analytics",
        "08_analytics_locked.jpeg",
        [("10-vote gate", "keeps early data honest."),
         ("Any mix counts", "video + image + TTS all add up."),
         ("Progress bar", "shows how close you are."),
         ("Then it opens", "full leaderboards appear.")],
        kick_color=YELLOW,
        img_box=(Inches(0.6), Inches(1.75), Inches(8.4), Inches(4.7)))

    # 12 Read results
    shot_full_slide(prs, D, nx(), "Read the results",
        "Analytics — who's winning, and where",
        "09_analytics_top.jpeg",
        [("Aggregate win rate", "head-to-head across all votes."),
         ("Per-mode & SKU", "T2V / I2V / R2V breakdowns."),
         ("Dimension radar", "strengths per metric."),
         ("Latency & voters", "speed and top evaluators.")],
        kick_color=GREEN,
        img_box=(Inches(0.6), Inches(1.7), Inches(8.5), Inches(5.05)))

    # 12b Win Rate by Tag
    shot_left_slide(prs, D, nx(), "Read the results",
        "Win Rate by Tag — where each model shines",
        "14_win_rate_by_tag.png",
        "Scroll down in Analytics for the tag matrix. Every case is auto-tagged (e.g. "
        "Photorealistic, People, Motion), so you can see which model wins for the content you "
        "care about — not just an overall average.",
        [("Rows = tags", "content categories, most-voted first."),
         ("Columns = models", "each competing model / variant."),
         ("Green vs red", "green = strong win rate, red = weak."),
         ("wins / total", "the sample size behind every cell.")],
        kick_color=GREEN)

    # 12c Evaluator leaderboard
    leaderboard_slide(prs, D, nx(), "Recognition",
        "The evaluator leaderboard — see where you rank",
        "Every vote you cast climbs the board. Rankings update live, per modality and overall — "
        "so you can see your standing against the rest of the team.",
        [("Ranked by votes", "the most active raters rise to the top.", BLUE),
         ("Per modality + overall", "video, image and TTS each have a board.", RED),
         ("Live & visible", "your LDAP and rank update instantly.", GREEN)],
        kick_color=YELLOW)

    # 13 Section: create your own
    section_slide(prs, "02", "Create your own evaluations",
        "Power users can generate fresh A/B cases from a simple JSON file — for any of the three modalities.")

    # 14 Admin overview
    shot_left_slide(prs, D, nx(), "Admin",
        "Generate new cases in the Admin console",
        "11_admin_image_gen.png",
        "In Admin, compose a single case in the form, or upload a cases JSON to batch-generate. "
        "Both models run automatically, an AI judge scores them, and the pair enters the blind queue.",
        [("Compose", "quick one-off from the form."),
         ("Upload JSON", "batch many cases at once."),
         ("Matchup registry", "pick which two models compete."),
         ("Generate + Auto-Eval", "runs both sides + AI judge.")],
        kick_color=BLUE)

    # 15 Video JSON
    code_slide(prs, D, nx(), "Input format · Video",
        "Video cases — a JSON array",
        "Upload on Admin → Video Generation. Each active video model matching the modality runs automatically.",
        '''[
  {
    "customer": "acme",
    "id": "V-1",
    "modality": "T2V",
    "prompt": "A green frog hops onto a lily pad, then a third.",
    "reference_images": [],
    "reference_videos": [],
    "aspect_ratio": "16:9",
    "duration": 8
  }
]''',
        [("prompt", "required — the generation prompt."),
         ("modality", "T2V · I2V · R2V."),
         ("id + customer", "dedup key (skips re-runs)."),
         ("reference_images", "required for I2V (first frame)."),
         ("reference_videos", "for R2V / V2V."),
         ("aspect_ratio / duration", "16:9 or 9:16 · 4–15s.")],
        kick_color=BLUE)

    # 16 Image JSON
    code_slide(prs, D, nx(), "Input format · Image",
        "Image cases — { \"cases\": [ … ] }",
        "Upload on Admin → Image Generation, with any reference images for i2i.",
        '''{
  "cases": [
    {
      "id": "img1",
      "mode": "t2i",
      "prompt": "A red maple leaf on white, studio photo",
      "matchup": "gemini-3.1-flash-image_vs_gpt2-medium"
    },
    {
      "id": "img3",
      "mode": "i2i",
      "prompt": "Make it snow; add a warm sunset glow",
      "input_image": "refs/koi.png",
      "matchup": "gemini-3-pro-image_vs_gpt2-high"
    }
  ]
}''',
        [("id", "required — unique case id."),
         ("mode", "t2i (text) or i2i (edit)."),
         ("prompt", "the generation / edit instruction."),
         ("input_image", "required for i2i."),
         ("matchup", "which model pair competes.")],
        kick_color=RED)

    # 17 TTS JSON
    code_slide(prs, D, nx(), "Input format · TTS",
        "TTS cases — { \"cases\": [ … ] }",
        "Upload on Admin → TTS Generation. Gemini TTS and ElevenLabs both synthesize each case.",
        '''{
  "cases": [
    {
      "id": "t1",
      "text": "[cheerful] Have a wonderful day!",
      "voice": "Kore",
      "style_prompt": "Warm, upbeat morning-radio host.",
      "language": "en"
    },
    {
      "id": "t3",
      "mode": "multi",
      "text": "Joe: How's the launch going?\\nJane: [excited] Better than we hoped!",
      "speakers": [
        { "speaker": "Joe",  "voice": "Kore" },
        { "speaker": "Jane", "voice": "Puck" }
      ]
    }
  ]
}''',
        [("text", "required — supports [tags] like [excited]."),
         ("voice", "a Gemini voice (Kore, Puck…)."),
         ("style_prompt", "delivery / tone direction."),
         ("language", "optional — auto-detected."),
         ("mode + speakers", "'multi' for 2-speaker chats.")],
        kick_color=GREEN)

    # 18 Closing
    s = blank(prs, INK); four_bar(s)
    text(s, Inches(0.9), Inches(1.7), Inches(11), Inches(0.5),
         "THAT'S IT", size=15, bold=True, color=RGBColor(0x9A,0xA0,0xA6), font=HEAD)
    text(s, Inches(0.88), Inches(2.2), Inches(11.5), Inches(1.6),
         [("Pick · Compare · Score · Decide", {"color": WHITE})],
         size=40, bold=True, font=HEAD, line_spc=1.05)
    _, tf = text(s, Inches(0.92), Inches(3.7), Inches(11), Inches(0.4),
                 [("Every honest vote sharpens the model choices we make for customers.",
                   {"color": RGBColor(0xBD,0xC1,0xC6), "size": 18})])
    para(tf, [("Cast your first 10 votes to unlock the leaderboards — then keep going.",
               {"color": RGBColor(0xBD,0xC1,0xC6), "size": 18})], space_before=8)
    para(tf, [("Questions or want a new evaluation set? Ping the Project Pulse admin.",
               {"color": RGBColor(0x9A,0xA0,0xA6), "size": 14})], space_before=18)
    for i, c in enumerate(FOUR):
        rect(s, Emu(int(Inches(0.94)) + i*int(Inches(0.34))), Inches(5.7),
             Inches(0.2), Inches(0.2), c, shape=MSO_SHAPE.OVAL)

    out = os.path.join(HERE, "GenMedia SxS - User Guide (How to Rate).pptx")
    prs.save(out)
    print("saved", out, "slides:", len(prs.slides._sldIdLst))
    return out


# ===========================================================================
# DECK 2 — LEADERSHIP / PRODUCT OVERVIEW
# ===========================================================================
def build_leadership():
    prs = new_deck()
    D = "GenMedia SxS · Project Pulse — Product Overview"
    n = [0]
    def nx(): n[0]+=1; return n[0]

    # 1 Title
    title_slide(
        prs, "Project Pulse · GenMedia SxS",
        [("Which model is\nactually ", {"color": G900}),
         ("best?", {"color": BLUE})],
        "One platform to prove — with blind human votes and an AI judge — which "
        "generative-media model wins for Video, Image and Speech.",
        "Product overview  ·  Internal benchmarking platform")

    # 2 Problem
    bullets_slide(prs, D, nx(), "The problem",
        "Model choice is exploding — opinions aren't proof",
        "New video, image and speech models ship every few weeks. Customers ask us which to use, "
        "and today the answer is often a gut feel.",
        [("Too many options", "Veo, Seedance, Kling, Gemini Image, GPT-image, ElevenLabs…", BLUE),
         ("Claims ≠ reality", "Benchmarks and demos rarely match a customer's real prompts.", RED),
         ("Slow & manual", "Ad-hoc comparisons in slides and folders don't scale.", YELLOW),
         ("Bias creeps in", "Knowing the model name skews the verdict.", GREEN)],
        two_col=True)

    # 3 Why it matters
    bullets_slide(prs, D, nx(), "Why it matters",
        "The stakes: our credibility and our consumption",
        "We advise APAC customers on Vertex AI generative media. Getting model choice right drives "
        "adoption, trust and token consumption.",
        [("Win customer POCs", "Show head-to-head proof on the customer's own prompts.", BLUE),
         ("Defensible advice", "Data, not opinion, behind every recommendation.", RED),
         ("Faster launches", "Benchmark a new model the day it lands.", YELLOW),
         ("Drives consumption", "Confident choices → more usage of the right models.", GREEN)],
        two_col=True)

    # 4 What it is
    s = blank(prs); four_bar(s)
    kicker(s, Inches(0.6), Inches(0.42), "The solution")
    text(s, Inches(0.6), Inches(0.78), Inches(12), Inches(0.7),
         "Project Pulse — one home for side-by-side evaluation", size=26, bold=True,
         color=G900, font=HEAD)
    text(s, Inches(0.6), Inches(1.55), Inches(12), Inches(0.6),
         "Two models, same prompt, names hidden. Humans vote and an AI judge scores. "
         "The result is a live leaderboard.", size=15.5, color=G700, font=BODY, line_spc=1.15)
    cards = [
        ("Blind & fair", "Model names hidden during voting to kill bias.", BLUE),
        ("3 modalities", "Video, Image and Speech — one workflow.", RED),
        ("Human + AI", "Real votes plus a Gemini judge on 5 metrics each.", YELLOW),
        ("Live leaderboard", "Win rates, radar and latency, in real time.", GREEN),
    ]
    x = Inches(0.6); cw = Inches(2.95); gap = Inches(0.11)
    for (h,d,c) in cards:
        card = rect(s, x, Inches(2.55), cw, Inches(2.5), G050, line=G200, line_w=Pt(0.75),
                    shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        try: card.adjustments[0]=0.05
        except Exception: pass
        rect(s, x, Inches(2.55), cw, Inches(0.1), c)
        text(s, Emu(int(x)+Emu(0.22*914400)), Inches(2.85),
             Emu(int(cw)-Emu(0.44*914400)), Inches(0.6), h, size=17, bold=True,
             color=G900, font=HEAD)
        text(s, Emu(int(x)+Emu(0.22*914400)), Inches(3.5),
             Emu(int(cw)-Emu(0.44*914400)), Inches(1.4), d, size=13.5, color=G700,
             font=BODY, line_spc=1.15)
        x = Emu(int(x)+int(cw)+int(gap))
    stat_row(s, Inches(5.35), [("3", "Modalities"), ("42", "Models tracked"),
                               ("2", "Signals: human + AI"), ("1", "Platform")])
    footer(s, D, nx())

    # 5 Product in action — arena
    shot_left_slide(prs, D, nx(), "In action",
        "Blind, side-by-side — the way people already judge",
        "06_video_blind_vote.png",
        "The core experience: two outputs for one prompt, model names hidden, scored on clear metrics. "
        "Simple enough that anyone can contribute a vote in seconds.",
        [("Same prompt", "both models get an identical brief."),
         ("Names hidden", "A vs B — no brand bias."),
         ("5 clear metrics", "structured, comparable scores."),
         ("Winner + why", "a verdict and a reason, every time.")],
        kick_color=RED)

    # 6 Two signals — AI evals
    shot_full_slide(prs, D, nx(), "Two independent signals",
        "Human votes + an AI judge — cross-checked",
        "07_ai_evals.jpeg",
        [("Human blind votes", "the ground truth of preference."),
         ("AI judge", "Gemini scores every pair on the same metrics."),
         ("Agreement = confidence", "when both agree, we're sure."),
         ("Scale", "the AI judge covers far more cases than humans alone.")],
        kick_color=BLUE,
        img_box=(Inches(0.6), Inches(1.75), Inches(8.4), Inches(4.7)))

    # 7 Insights dashboard
    shot_full_slide(prs, D, nx(), "Insights that drive decisions",
        "From votes to a defensible recommendation",
        "09_analytics_top.jpeg",
        [("Head-to-head win rates", "who wins, overall and per mode."),
         ("Per-metric radar", "where each model is strong or weak."),
         ("Latency", "quality vs speed trade-offs."),
         ("Filter by category", "answers for a customer's use-case.")],
        kick_color=GREEN,
        img_box=(Inches(0.6), Inches(1.7), Inches(8.5), Inches(5.05)))

    # 7b Win rate by tag
    shot_left_slide(prs, D, nx(), "Insights",
        "Win rates by use-case — not just an average",
        "14_win_rate_by_tag.png",
        "The tag matrix breaks results down by content type, so we can recommend the best model "
        "for a customer's actual use-case rather than a single blended number.",
        [("Per use-case truth", "best model for Food, People, Motion…"),
         ("Strengths & gaps", "green / red shows where each model leads."),
         ("Backed by samples", "every cell shows wins / total."),
         ("Targeted advice", "match the model to the customer's content.")],
        kick_color=BLUE)

    # 7c Evaluator leaderboard
    leaderboard_slide(prs, D, nx(), "Engagement",
        "A live leaderboard — the engine for participation",
        "The evaluator leaderboard turns rating into friendly competition — the natural hook for a "
        "recognition or spot-bonus program to scale votes and sharpen our model calls.",
        [("Recognition built-in", "top raters are visible to the whole team.", BLUE),
         ("Fuels more data", "competition drives volume → better calls.", GREEN),
         ("Ready for incentives", "rank = basis for a quarterly spot bonus.", YELLOW)],
        kick_color=YELLOW)

    # 8 Built for scale (stats)
    s = blank(prs); four_bar(s)
    kicker(s, Inches(0.6), Inches(0.42), "Scale")
    text(s, Inches(0.6), Inches(0.78), Inches(12), Inches(0.7),
         "Built to benchmark at volume", size=27, bold=True, color=G900, font=HEAD)
    text(s, Inches(0.6), Inches(1.5), Inches(12), Inches(0.6),
         "Automated generation + AI judging means coverage scales far beyond manual review.",
         size=15, color=G700, font=BODY)
    stat_row(s, Inches(2.3), [("282", "AI auto-evals"), ("42", "Models tracked"),
                              ("3", "Modalities"), ("5", "Metrics / modality")])
    stat_row(s, Inches(4.05), [("53", "Video evals"), ("199", "Image evals"),
                               ("30", "TTS evals"), ("100s", "Human votes")])
    text(s, Inches(0.6), Inches(5.95), Inches(12), Inches(0.6),
         "Models span Veo 3.1, Seedance, Kling & Grok Imagine (video), Gemini Image vs GPT-image, "
         "and Gemini TTS vs ElevenLabs.", size=13, color=G500, font=BODY, line_spc=1.15)
    footer(s, D, nx())

    # 9 What it solves
    bullets_slide(prs, D, nx(), "Outcomes",
        "What Project Pulse solves",
        None,
        [("Objective model selection", "Blind, structured scoring removes brand bias.", BLUE),
         ("Customer-specific proof", "Run the customer's own prompts and categories.", RED),
         ("Day-one benchmarking", "Score a new model as soon as it ships.", YELLOW),
         ("A shared source of truth", "One leaderboard the whole team can cite.", GREEN),
         ("Human + AI confidence", "Two signals that cross-check each other.", BLUE),
         ("Repeatable & auditable", "Every case, score and winner is recorded.", RED)],
        two_col=True)

    # 10 Use cases
    bullets_slide(prs, D, nx(), "Where it's used",
        "Who uses it, and for what",
        None,
        [("Sales & Customer Eng.", "Head-to-head proof inside a POC.", BLUE),
         ("Product & PMs", "Track model quality release over release.", RED),
         ("Model launches", "Independent read on a new model's strengths.", YELLOW),
         ("Localization", "Multilingual prompts — incl. mixed-language TTS.", GREEN)],
        two_col=True)

    # 11 Why it wins / differentiators
    bullets_slide(prs, D, nx(), "Why this approach wins",
        "What makes it different",
        None,
        [("Blind by design", "Bias removed where verdicts are formed.", BLUE),
         ("Multimodal in one place", "Video, image and speech share one workflow.", RED),
         ("Human + AI together", "Preference and scale, cross-validated.", YELLOW),
         ("Self-serve evals", "Anyone can add cases via a simple JSON.", GREEN)],
        two_col=True)

    # 12 Close / CTA
    s = blank(prs, INK); four_bar(s)
    text(s, Inches(0.9), Inches(1.7), Inches(11), Inches(0.5),
         "THE ASK", size=15, bold=True, color=RGBColor(0x9A,0xA0,0xA6), font=HEAD)
    text(s, Inches(0.88), Inches(2.2), Inches(11.5), Inches(1.6),
         [("Make model choice ", {"color": WHITE}), ("evidence-based.", {"color": BLUE})],
         size=40, bold=True, font=HEAD, line_spc=1.05)
    _, tf = text(s, Inches(0.92), Inches(3.7), Inches(11.2), Inches(0.4),
                 [("Use Project Pulse in your next POC to show head-to-head proof on the customer's prompts.",
                   {"color": RGBColor(0xBD,0xC1,0xC6), "size": 18})])
    para(tf, [("Contribute blind votes so the leaderboard reflects real preference.",
               {"color": RGBColor(0xBD,0xC1,0xC6), "size": 18})], space_before=10)
    para(tf, [("Live · genmedia-sxs-v2 (Google-IAM access).",
               {"color": RGBColor(0x9A,0xA0,0xA6), "size": 14})], space_before=18)
    for i, c in enumerate(FOUR):
        rect(s, Emu(int(Inches(0.94)) + i*int(Inches(0.34))), Inches(5.7),
             Inches(0.2), Inches(0.2), c, shape=MSO_SHAPE.OVAL)

    out = os.path.join(HERE, "GenMedia SxS - Product Overview (Leadership).pptx")
    prs.save(out)
    print("saved", out, "slides:", len(prs.slides._sldIdLst))
    return out


if __name__ == "__main__":
    build_user_guide()
    build_leadership()

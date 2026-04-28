#!/usr/bin/env python3
"""Generate GenMedia SxS Evaluator Guidelines slide deck in Google brand style."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

ASSETS = os.path.join(os.path.dirname(__file__), "assets")

# Google brand colors
GOOGLE_BLUE = RGBColor(0x42, 0x85, 0xF4)
GOOGLE_RED = RGBColor(0xDB, 0x44, 0x37)
GOOGLE_YELLOW = RGBColor(0xF4, 0xB4, 0x00)
GOOGLE_GREEN = RGBColor(0x0F, 0x9D, 0x58)
GOOGLE_GREY_900 = RGBColor(0x20, 0x21, 0x24)
GOOGLE_GREY_700 = RGBColor(0x5F, 0x63, 0x68)
GOOGLE_GREY_200 = RGBColor(0xE8, 0xEA, 0xED)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def set_slide_bg(slide, color):
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_colored_bar(slide, left, top, width, height, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 bold=False, color=GOOGLE_GREY_900, alignment=PP_ALIGN.LEFT,
                 font_name="Google Sans"):
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = font_name
    p.alignment = alignment
    return tf


def add_bullet_slide_content(tf, items, font_size=16, color=GOOGLE_GREY_700, font_name="Google Sans"):
    for item in items:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = font_name
        p.space_before = Pt(6)
        p.level = 0


def add_google_dots(slide, top):
    colors = [GOOGLE_BLUE, GOOGLE_RED, GOOGLE_YELLOW, GOOGLE_GREEN]
    start_left = Inches(6.0)
    for i, c in enumerate(colors):
        dot = slide.shapes.add_shape(
            MSO_SHAPE.OVAL,
            start_left + Inches(i * 0.35), top,
            Inches(0.18), Inches(0.18)
        )
        dot.fill.solid()
        dot.fill.fore_color.rgb = c
        dot.line.fill.background()


def build():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    # =========================================================
    # SLIDE 1 — Title
    # =========================================================
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, WHITE)
    add_colored_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.08), GOOGLE_BLUE)

    add_text_box(slide, Inches(1), Inches(1.8), Inches(11), Inches(1.2),
                 "GenMedia SxS", font_size=52, bold=True, color=GOOGLE_BLUE)
    add_text_box(slide, Inches(1), Inches(3.0), Inches(11), Inches(0.8),
                 "Evaluator Guidelines", font_size=36, color=GOOGLE_GREY_900)
    add_text_box(slide, Inches(1), Inches(4.2), Inches(11), Inches(0.6),
                 "Side-by-side blind evaluation of generative media models",
                 font_size=18, color=GOOGLE_GREY_700)

    add_google_dots(slide, Inches(5.5))

    add_text_box(slide, Inches(1), Inches(6.4), Inches(5), Inches(0.4),
                 "Google  |  Internal Only", font_size=14, color=GOOGLE_GREY_700)

    # =========================================================
    # SLIDE 2 — Access the Arena
    # =========================================================
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, WHITE)
    add_colored_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOOGLE_RED)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(5), Inches(0.6),
                 "1. Access the Arena", font_size=32, bold=True, color=GOOGLE_GREY_900)

    tf = add_text_box(slide, Inches(0.8), Inches(1.2), Inches(4.5), Inches(3),
                      "", font_size=16, color=GOOGLE_GREY_700)
    tf.paragraphs[0].text = "Getting started"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = GOOGLE_GREY_900
    add_bullet_slide_content(tf, [
        "Navigate to the platform URL",
        "Enter your LDAP handle in the login field",
        "Press 'Enter Evaluation' to begin",
        "The leaderboard shows top evaluators in real-time",
        "Your vote count is tracked across sessions",
    ])

    img_path = os.path.join(ASSETS, "01_login_leaderboard.png")
    slide.shapes.add_picture(img_path, Inches(5.8), Inches(1.0), Inches(7.0))

    # =========================================================
    # SLIDE 3 — SxS Evaluation
    # =========================================================
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, WHITE)
    add_colored_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOOGLE_BLUE)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
                 "2. Side-by-Side Evaluation", font_size=32, bold=True, color=GOOGLE_GREY_900)

    img_path = os.path.join(ASSETS, "02_evaluation_arena.png")
    slide.shapes.add_picture(img_path, Inches(0.8), Inches(1.3), Inches(11.7))

    add_text_box(slide, Inches(0.8), Inches(6.2), Inches(11), Inches(1),
                 "Two anonymized video variants from the same prompt. Model identities are hidden until after you vote.",
                 font_size=16, color=GOOGLE_GREY_700)

    # =========================================================
    # SLIDE 4 — Voting Flow
    # =========================================================
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, WHITE)
    add_colored_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOOGLE_GREEN)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(5), Inches(0.6),
                 "3. Voting Flow", font_size=32, bold=True, color=GOOGLE_GREY_900)

    tf = add_text_box(slide, Inches(0.8), Inches(1.2), Inches(4.5), Inches(4),
                      "", font_size=16, color=GOOGLE_GREY_700)
    tf.paragraphs[0].text = "How to vote"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = GOOGLE_GREY_900
    add_bullet_slide_content(tf, [
        "Watch both variants at least twice",
        "Use 'Play Both' for synchronized playback",
        "Click 'Select Variant A' or 'Select Variant B'",
        "Complete the detailed scoring rubric",
        "Use 'Skip' if neither variant is evaluable",
    ])

    img_path = os.path.join(ASSETS, "03_evaluation_full.png")
    slide.shapes.add_picture(img_path, Inches(5.8), Inches(0.9), Inches(7.0))

    # =========================================================
    # SLIDE 5 — Evaluation Criteria
    # =========================================================
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, WHITE)
    add_colored_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOOGLE_YELLOW)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
                 "4. Evaluation Criteria", font_size=32, bold=True, color=GOOGLE_GREY_900)

    criteria = [
        ("Prompt Adherence", "Does the video faithfully execute all stated elements, timings, and visual requests?", GOOGLE_BLUE),
        ("Motion Quality", "Is the movement natural and fluid? Any temporal morphing or flickering artifacts?", GOOGLE_RED),
        ("Aesthetic Quality", "Assess lighting, resolution, color grading, and cinematic quality.", GOOGLE_YELLOW),
        ("Physics & Cohesion", "Do objects interact realistically? Gravity, collisions, spatial consistency?", GOOGLE_GREEN),
        ("Audio Expressiveness", "If audio is present, does it match the scene's mood and energy?", GOOGLE_BLUE),
        ("Audio-Visual Sync", "Are sounds timed correctly with on-screen actions?", GOOGLE_RED),
    ]

    for i, (title, desc, accent) in enumerate(criteria):
        row = i // 2
        col = i % 2
        x = Inches(0.8 + col * 6.2)
        y = Inches(1.3 + row * 1.8)

        add_colored_bar(slide, x, y, Inches(0.08), Inches(1.2), accent)

        add_text_box(slide, x + Inches(0.3), y + Inches(0.05), Inches(5.5), Inches(0.4),
                     title, font_size=20, bold=True, color=GOOGLE_GREY_900)
        add_text_box(slide, x + Inches(0.3), y + Inches(0.5), Inches(5.5), Inches(0.7),
                     desc, font_size=15, color=GOOGLE_GREY_700)

    # =========================================================
    # SLIDE 6 — The 10-Vote Gate & Analytics
    # =========================================================
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, WHITE)
    add_colored_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOOGLE_GREEN)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(5), Inches(0.6),
                 "5. The 10-Vote Gate", font_size=32, bold=True, color=GOOGLE_GREY_900)

    tf = add_text_box(slide, Inches(0.8), Inches(1.2), Inches(4.5), Inches(4),
                      "", font_size=16, color=GOOGLE_GREY_700)
    tf.paragraphs[0].text = "Unlock analytics by evaluating"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = GOOGLE_GREY_900
    add_bullet_slide_content(tf, [
        "A progress bar in the top-right tracks your votes",
        "Complete 10 evaluations to unlock the Analytics button",
        "Once unlocked, access the full Benchmark Results dashboard",
        "You can also submit custom benchmarking prompts",
        "The leaderboard ranks all evaluators by vote count",
    ])

    img_path = os.path.join(ASSETS, "08_analytics_unlocked.png")
    slide.shapes.add_picture(img_path, Inches(5.8), Inches(1.0), Inches(7.0))

    # =========================================================
    # SLIDE 7 — Benchmark Results
    # =========================================================
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, WHITE)
    add_colored_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOOGLE_BLUE)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
                 "6. Benchmark Results", font_size=32, bold=True, color=GOOGLE_GREY_900)
    add_text_box(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.5),
                 "Aggregated win rates, mode breakup, per-SKU performance with latency, and radar charts across evaluation dimensions.",
                 font_size=16, color=GOOGLE_GREY_700)

    img_path = os.path.join(ASSETS, "09_analytics_dashboard.png")
    slide.shapes.add_picture(img_path, Inches(1.5), Inches(1.9), Inches(10.3))

    # =========================================================
    # SLIDE 8 — Leaderboard & Full Analytics
    # =========================================================
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, WHITE)
    add_colored_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOOGLE_YELLOW)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(5), Inches(0.6),
                 "7. Leaderboard", font_size=32, bold=True, color=GOOGLE_GREY_900)

    tf = add_text_box(slide, Inches(0.8), Inches(1.2), Inches(4.5), Inches(4.5),
                      "", font_size=16, color=GOOGLE_GREY_700)
    tf.paragraphs[0].text = "Gamified participation"
    tf.paragraphs[0].font.size = Pt(20)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = GOOGLE_GREY_900
    add_bullet_slide_content(tf, [
        "Top Evaluators ranked by total vote count",
        "Visible on login screen and analytics dashboard",
        "Dimension Analytics with radar charts for T2V and I2V",
        "Toggle between Global and Personal results",
        "Filter by prompt category tags",
        "'Evaluation Tier Unlocked' badge at 10 votes",
    ])

    img_path = os.path.join(ASSETS, "10_analytics_full.png")
    slide.shapes.add_picture(img_path, Inches(5.8), Inches(0.8), height=Inches(6.2))

    # =========================================================
    # SLIDE 9 — Tips
    # =========================================================
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, WHITE)
    add_colored_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOOGLE_YELLOW)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
                 "8. Tips for Great Evaluations", font_size=32, bold=True, color=GOOGLE_GREY_900)

    tips = [
        ("Watch each video at least twice", "First pass for overall impression, second for details and artifacts.", GOOGLE_BLUE),
        ("Use audio when available", "Audio expressiveness and sync are scoring dimensions.", GOOGLE_RED),
        ("Don't rush", "Quality evaluations are more valuable than quantity.", GOOGLE_YELLOW),
        ("Skip freely", "If a pair is ambiguous or both outputs are equally bad/good, skip it.", GOOGLE_GREEN),
    ]

    for i, (title, desc, accent) in enumerate(tips):
        y = Inches(1.4 + i * 1.35)
        add_colored_bar(slide, Inches(1.5), y, Inches(0.08), Inches(0.9), accent)
        add_text_box(slide, Inches(1.9), y + Inches(0.0), Inches(9), Inches(0.4),
                     title, font_size=22, bold=True, color=GOOGLE_GREY_900)
        add_text_box(slide, Inches(1.9), y + Inches(0.45), Inches(9), Inches(0.5),
                     desc, font_size=16, color=GOOGLE_GREY_700)

    add_google_dots(slide, Inches(6.8))

    # =========================================================
    # SLIDE 10 — Admin: Prompt Engine
    # =========================================================
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, WHITE)
    add_colored_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOOGLE_BLUE)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
                 "9. Admin — Prompt Engine", font_size=32, bold=True, color=GOOGLE_GREY_900)
    add_text_box(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.5),
                 "Create Hard Prompts with T2V, I2V, and R2V modes. Upload start/end frames, assign category tags, and trigger generation across all active models.",
                 font_size=16, color=GOOGLE_GREY_700)

    img_path = os.path.join(ASSETS, "05_admin_prompts.png")
    slide.shapes.add_picture(img_path, Inches(1.5), Inches(1.9), Inches(10.3))

    # =========================================================
    # SLIDE 11 — Admin: Generations
    # =========================================================
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, WHITE)
    add_colored_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOOGLE_GREEN)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
                 "10. Admin — Generations Review", font_size=32, bold=True, color=GOOGLE_GREY_900)
    add_text_box(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.5),
                 "Track all generation jobs with status indicators, category tags, and success/failure counts. Filter by status, tags, or search by prompt text.",
                 font_size=16, color=GOOGLE_GREY_700)

    img_path = os.path.join(ASSETS, "06_admin_generations.png")
    slide.shapes.add_picture(img_path, Inches(1.5), Inches(1.9), Inches(10.3))

    # =========================================================
    # SLIDE 12 — Admin: Model Registry
    # =========================================================
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, WHITE)
    add_colored_bar(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.06), GOOGLE_RED)

    add_text_box(slide, Inches(0.8), Inches(0.4), Inches(11), Inches(0.6),
                 "11. Admin — Model Registry", font_size=32, bold=True, color=GOOGLE_GREY_900)
    add_text_box(slide, Inches(0.8), Inches(1.1), Inches(11), Inches(0.5),
                 "View and toggle all registered models. Supports Veo (Vertex AI), Kling, Seedance, and other providers.",
                 font_size=16, color=GOOGLE_GREY_700)

    img_path = os.path.join(ASSETS, "07_admin_models.png")
    slide.shapes.add_picture(img_path, Inches(1.5), Inches(1.9), Inches(10.3))

    # =========================================================
    # SLIDE 13 — Thank You
    # =========================================================
    slide = prs.slides.add_slide(blank)
    set_slide_bg(slide, GOOGLE_BLUE)

    add_text_box(slide, Inches(1), Inches(2.5), Inches(11), Inches(1),
                 "Thank you!", font_size=52, bold=True, color=WHITE,
                 alignment=PP_ALIGN.CENTER)
    add_text_box(slide, Inches(1), Inches(3.8), Inches(11), Inches(0.8),
                 "Your evaluations directly shape the quality of our generative media models.",
                 font_size=20, color=WHITE, alignment=PP_ALIGN.CENTER)
    add_text_box(slide, Inches(1), Inches(5.5), Inches(11), Inches(0.5),
                 "Complete 10 evaluations to unlock prompt submission",
                 font_size=16, color=RGBColor(0xA0, 0xC4, 0xF8), alignment=PP_ALIGN.CENTER)

    # Save
    out = os.path.join(os.path.dirname(__file__), "GenMedia_SxS_Evaluator_Guidelines.pptx")
    prs.save(out)
    print(f"Saved: {out}")


if __name__ == "__main__":
    build()
